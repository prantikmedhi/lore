"""Create a PPTX deck from a Markdown research brief or JSON outline."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path


def read_input(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def markdown_sections(text: str) -> list[tuple[str, list[str]]]:
    sections: list[tuple[str, list[str]]] = []
    current_title = "Overview"
    current_items: list[str] = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        heading = re.match(r"^#{1,3}\s+(.+)$", line)
        if heading:
            if current_items:
                sections.append((current_title, current_items))
            current_title = heading.group(1)
            current_items = []
            continue
        if line.startswith("- "):
            current_items.append(line[2:])
        elif len(line) > 20:
            current_items.append(line)
    if current_items:
        sections.append((current_title, current_items))
    return sections or [("Overview", [text[:400]])]


def load_sections(path: Path) -> tuple[str, list[tuple[str, list[str]]]]:
    text = read_input(path)
    if path.suffix.lower() == ".json":
        data = json.loads(text)
        title = data.get("title") or data.get("workflow") or "Lore Deck"
        slides = data.get("slides") or data.get("deliverables") or data.get("steps") or []
        sections = [(str(item), []) if not isinstance(item, dict) else (str(item.get("title", "Slide")), item.get("bullets", [])) for item in slides]
        return title, sections or [("Overview", ["No slide content provided."])]
    sections = markdown_sections(text)
    return sections[0][0], sections


def add_slide(prs, title: str, bullets: list[str]) -> None:
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title[:90]
    body = slide.placeholders[1].text_frame
    body.clear()
    for idx, bullet in enumerate(bullets[:6] or ["Add source-grounded content here."]):
        paragraph = body.paragraphs[0] if idx == 0 else body.add_paragraph()
        paragraph.text = bullet[:240]
        paragraph.level = 0


def write_pptx(input_path: Path, output_path: Path, max_slides: int) -> None:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise SystemExit("python-pptx is required. Install with: pip install python-pptx") from exc

    title, sections = load_sections(input_path)
    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = title[:90]
    title_slide.placeholders[1].text = "Lore | en source-grounded research"
    for section_title, bullets in sections[: max(1, max_slides - 1)]:
        add_slide(prs, section_title, bullets)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="lore-ppt")
    parser.add_argument("--input", required=True, help="Markdown brief or JSON outline.")
    parser.add_argument("--output", required=True, help="Output .pptx path.")
    parser.add_argument("--max-slides", type=int, default=10)
    return parser


def main() -> int:
    ns = build_parser().parse_args()
    write_pptx(Path(ns.input), Path(ns.output), ns.max_slides)
    print(json.dumps({"created": ns.output, "language": "en"}, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
